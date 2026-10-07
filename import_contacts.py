"""Import teams and contacts from the league spreadsheet.

Usage: python3 import_contacts.py "/path/to/Tuesday Doubles Team League.xlsx"

Reads the "Teams & Boxes" sheet (a "Team <captain>" header row followed by that
team's players) and the "Email List" sheet. Writes data/teams.json and
data/contacts.json (both gitignored - real names and emails) and adds any team
member missing from data/players.json at the default rating. Safe to re-run.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl

from league_logic import DEFAULT_RATING, match_contact_to_roster

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"


def _write_json(path: Path, data: object) -> None:
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(data, indent=2), encoding="utf-8")
    tmp.replace(path)


def load_roster() -> list[dict]:
    raw = json.loads((DATA_DIR / "players.json").read_text(encoding="utf-8"))
    return [
        {"name": p, "rating": DEFAULT_RATING} if isinstance(p, str) else p
        for p in raw
    ]


def import_teams(sheet, players: list[dict]) -> dict[str, list[str]]:
    """Parse team blocks; the header row names the captain, who also plays."""
    names = [p["name"] for p in players]

    def resolve(raw_name: str) -> str:
        match = match_contact_to_roster(raw_name, names)
        if match:
            return match
        names.append(raw_name)
        players.append({"name": raw_name, "rating": DEFAULT_RATING})
        return raw_name

    teams: dict[str, list[str]] = {}
    current: str | None = None
    for (cell,) in sheet.iter_rows(min_col=1, max_col=1, values_only=True):
        if not cell or not str(cell).strip():
            continue
        text = str(cell).strip()
        if text.lower().startswith("team "):
            current = text[5:].strip()
            captain = match_contact_to_roster(text, names) or current
            teams[current] = [resolve(captain)]
        elif current:
            member = resolve(text)
            if member not in teams[current]:
                teams[current].append(member)
    return teams


def main(xlsx_path: str) -> None:
    workbook = openpyxl.load_workbook(xlsx_path, data_only=True)
    DATA_DIR.mkdir(exist_ok=True)

    players = load_roster()
    teams = import_teams(workbook["Teams & Boxes"], players)
    _write_json(DATA_DIR / "teams.json", teams)
    _write_json(DATA_DIR / "players.json", players)
    roster = [p["name"] for p in players]

    contacts = []
    for name, email, *_ in workbook["Email List"].iter_rows(min_row=2, values_only=True):
        if not name or not email:
            continue
        contacts.append({
            "name": str(name).strip(),
            "email": str(email).strip(),
            "roster_name": match_contact_to_roster(str(name), roster),
        })
    _write_json(DATA_DIR / "contacts.json", contacts)

    on_teams = {p for members in teams.values() for p in members}
    print(f"Teams: {', '.join(f'{t} ({len(m)})' for t, m in teams.items())}")
    print(f"Contacts: {len(contacts)} ({sum(1 for c in contacts if c['roster_name'])} matched to the roster)")
    print(f"On the roster but on no team: {[n for n in roster if n not in on_teams]}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
